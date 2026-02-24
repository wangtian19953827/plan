#!/usr/bin/env python3
"""
创建示例 Excel 文件用于测试导入功能
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

def create_sample_excel(filename='sample_data.xlsx'):
    """创建示例 Excel 文件"""

    # 创建工作簿
    wb = openpyxl.Workbook()

    # ========== Sheet 1: 产品信息 ==========
    ws1 = wb.active
    ws1.title = "产品信息"

    # 表头
    headers1 = ["产品编码", "产品名称", "单价", "描述"]
    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 示例产品数据
    products = [
        ("P001", "智能手机 Pro", 5999, "旗舰智能手机，128GB存储"),
        ("P002", "平板电脑 Air", 3999, "10.9英寸平板，WiFi版"),
        ("P003", "无线耳机 Pro", 1499, "主动降噪无线耳机"),
        ("P004", "智能手表", 2299, "健康监测智能手表"),
        ("P005", "超极本", 8999, "轻薄超极本，16GB内存"),
        ("P006", "机械键盘", 799, "RGB背光机械键盘"),
        ("P007", "游戏鼠标", 599, "高精度游戏鼠标"),
        ("P008", "显示器 4K", 3999, "27英寸4K显示器"),
        ("P009", "移动电源", 299, "20000mAh移动电源"),
        ("P010", "充电器", 199, "65W快充充电器"),
    ]

    # 添加更多产品到80个
    for i in range(11, 81):
        code = f"P{i:03d}"
        products.append((
            code,
            f"通用产品{i}",
            100 + i * 10,
            f"标准产品{i}描述"
        ))

    # 写入产品数据
    for row_idx, product in enumerate(products, 2):
        for col_idx, value in enumerate(product, 1):
            ws1.cell(row=row_idx, column=col_idx, value=value)

    # 设置列宽
    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 20
    ws1.column_dimensions['C'].width = 10
    ws1.column_dimensions['D'].width = 30

    # ========== Sheet 2: 月度数据 ==========
    ws2 = wb.create_sheet("月度数据")

    # 表头
    headers2 = ["产品编码", "年份", "月份", "月度需求", "月末目标库存", "月初库存"]
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 月度数据
    current_year = 2025
    current_month = 2

    for row_idx, product in enumerate(products, 2):
        product_code = product[0]
        # 根据产品编号生成不同的数据
        demand = 50 + (row_idx % 100) * 5
        end_stock = 30 + (row_idx % 50) * 4
        start_stock = 50 + (row_idx % 60) * 3

        ws2.cell(row=row_idx, column=1, value=product_code)
        ws2.cell(row=row_idx, column=2, value=current_year)
        ws2.cell(row=row_idx, column=3, value=current_month)
        ws2.cell(row=row_idx, column=4, value=demand)
        ws2.cell(row=row_idx, column=5, value=end_stock)
        ws2.cell(row=row_idx, column=6, value=start_stock)

    # 设置列宽
    ws2.column_dimensions['A'].width = 12
    ws2.column_dimensions['B'].width = 8
    ws2.column_dimensions['C'].width = 6
    ws2.column_dimensions['D'].width = 12
    ws2.column_dimensions['E'].width = 14
    ws2.column_dimensions['F'].width = 12

    # 保存文件
    wb.save(filename)
    print(f"✓ 已创建示例文件: {filename}")
    print(f"  - 产品数量: {len(products)}")
    print(f"  - 工作表数量: 2")
    print(f"  - 文件大小: {len(products)} 行产品数据 + {len(products)} 行月度数据")

if __name__ == "__main__":
    try:
        create_sample_excel()
        print("\n现在可以在程序中点击'导入Excel'并选择 sample_data.xlsx 进行测试！")
    except ImportError:
        print("错误: 需要安装 openpyxl 库")
        print("运行: pip install openpyxl")
    except Exception as e:
        print(f"错误: {e}")
